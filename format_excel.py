from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

# Excel file location 
file_path = Path(__file__).resolve().parent / 'books_catalog.xlsx'

# Open Excel workbook
workbook = load_workbook(file_path)

# Select active worksheet
worksheet = workbook.active

# Formatar header
for cell in worksheet[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(
        fill_type= 'solid',
        fgColor='0070C0'
    )

# Save formatted workbook
workbook.save(file_path)

print('Excel formatting completed!')